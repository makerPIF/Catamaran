#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
대전중앙중학교 USV - BOM / 예산 / 추진성능 계산 워크북 생성 스크립트

사용법:
    pip install openpyxl
    python build_bom.py
    -> 대전중앙중_USV_BOM.xlsx 생성

워크북 구성
    1) BOM          전체 부품표 (소계는 수식)
    2) 예산요약      구분별 합계 (SUMIF)
    3) 추진성능계산  A2212 적합성 검증 계산기 (노란 칸을 바꾸면 자동 재계산)
    4) 통신방식비교  장거리 통신 확보 방안
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = "맑은 고딕"
BLUE = Font(name=FONT, size=10, color="0000FF")          # 사용자가 바꾸는 입력값
BLACK = Font(name=FONT, size=10)                         # 계산/일반
BOLD = Font(name=FONT, size=10, bold=True)
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=9, color="666666")
HFILL = PatternFill("solid", fgColor="2A78D6")
GFILL = PatternFill("solid", fgColor="EDF3FC")
YFILL = PatternFill("solid", fgColor="FFF3CD")
thin = Side(style="thin", color="D5D5D5")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# =====================================================================
# 시트 1 : BOM
# =====================================================================
ws = wb.active
ws.title = "BOM"

ws["A1"] = "대전중앙중학교 3D프린팅 쌍동선 무인정(USV) 부품표 (BOM)"
ws["A1"].font = TITLE
ws["A2"] = ("단가는 2026년 국내 온라인 유통 기준 참고가(원)입니다. 실제 구매 전 반드시 재확인하세요. "
            "'선택' 표기 부품은 없어도 기본 동작합니다.")
ws["A2"].font = SUB
ws.merge_cells("A1:H1")
ws.merge_cells("A2:H2")

HEADERS = ["번호", "구분", "부품명", "규격 / 추천 모델", "수량", "단가(원)", "소계(원)", "비고"]
for i, h in enumerate(HEADERS, 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = HEAD
    c.fill = HFILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")

# (구분, 부품명, 규격, 수량, 단가, 비고)
ITEMS = [
    ("A. 추진계", "BLDC 모터 A2212 10T", "1400KV / 2–3S / 최대 180W·14A / 축 3.17mm", 2, 13000,
     "★ 에어프로펠러(수면 위) 방식 전용. 물에 잠기면 안 됨"),
    ("A. 추진계", "ESC 30A (BLHeli_S)", "2–4S, BEC 5V 2A, BLHeli_S 펌웨어", 2, 18000,
     "★ 반드시 BLHeli_S여야 3D(후진) 모드 설정 가능"),
    ("A. 추진계", "프로펠러 8045 CW/CCW 세트", '8×4.5" 나일론+GF, 허브 3.17mm', 3, 4000,
     "1세트=CW1+CCW1. 파손 대비 여유분 포함"),
    ("A. 추진계", "BLHeli USB 링커(프로그래밍 카드)", "SILABS C2 / 1-wire 지원", 1, 12000,
     "ESC 3D모드·회전방향 설정용. 팀당 1개면 충분"),
    ("A. 추진계", "3.5mm 바나나 커넥터 세트", "금도금 암/수", 6, 800, "모터–ESC 연결"),
    ("A. 추진계", "모터 마운트 (3D 프린팅)", "PETG, 16mm 볼트 패턴", 2, 0, "자체 제작 · 제작 가이드 참조"),
    ("A. 추진계", "프로펠러 안전 가드 (3D 프린팅)", "내경 Ø225mm 링 + 메쉬", 2, 0, "★ 안전상 필수. 손가락 보호"),

    ("B. 제어·센서", "ESP32 NodeMCU DevKit V1", "ESP32-WROOM-32, 30핀", 1, 12000, "메인 컨트롤러"),
    ("B. 제어·센서", "MPU6050 (GY-521)", "6축 IMU, I2C, 3.3–5V", 1, 3000, "자세(Roll/Pitch/Yaw)"),
    ("B. 제어·센서", "u-blox NEO-8M GPS 모듈", "세라믹 안테나 포함, UART 9600bps", 1, 22000, "위치·속도·침로"),
    ("B. 제어·센서", "XIAO ESP32S3 Sense", "OV2640 카메라 + PSRAM 8MB 포함", 1, 32000, "영상 송신 전용"),
    ("B. 제어·센서", "능동 부저 5V", "3.3V 구동 가능형", 1, 1000, "저전압·아밍 경보"),
    ("B. 제어·센서", "저항 100kΩ / 22kΩ (1% 금속피막)", "1/4W", 4, 200, "배터리 전압 분배 회로"),
    ("B. 제어·센서", "만능기판 + 핀헤더 + 점퍼선 세트", "5×7cm 기판 포함", 1, 8000, "배선 정리용"),
    ("B. 제어·센서", "XT60 커넥터 암/수", "고전류 배터리 커넥터", 3, 2500, ""),

    ("C. 전원", "LiPo 배터리 3S 5000mAh 30C", "11.1V, XT60 단자", 2, 55000, "1개는 예비. ★ 3S 고정"),
    ("C. 전원", "UBEC 5V 3A (스위칭 방식)", "입력 2–6S", 2, 9000, "①ESP32계 ②공유기 전원 분리"),
    ("C. 전원", "LiPo 밸런스 충전기", "IMAX B6 급 + 어댑터", 1, 45000, "학교 보유 시 제외"),
    ("C. 전원", "방수 토글 스위치", "12V 20A 이상", 1, 4000, "메인 전원 차단"),
    ("C. 전원", "XT60 병렬 분배 하네스", "1→3 분기", 1, 6000, "ESC 2개 + UBEC"),
    ("C. 전원", "LiPo 안전 보관 가방", "방염", 1, 8000, "★ 안전상 필수"),
    ("C. 전원", "배터리 잔량 체커/경보기", "1–8S 셀 전압 표시", 1, 4000, ""),

    ("D. 통신", "미니 공유기 GL.iNet GL-AR300M16-Ext", "OpenWrt / 외장안테나 2 / 5V 입력", 1, 45000,
     "USV 탑재. SSID: 대전중앙중학교"),
    ("D. 통신", "2.4GHz 5dBi 무지향 안테나 (RP-SMA)", "공유기 기본 안테나 대체", 2, 6000, "USV 측"),
    ("D. 통신", "지상국 USB 무선랜 (외부안테나형)", "ALFA AWUS036NHA 급, RP-SMA", 1, 55000,
     "★ 장거리 확보의 핵심"),
    ("D. 통신", "2.4GHz 14dBi 지향성 패널 안테나", "RP-SMA, 실외형", 1, 35000, "★ 지상국에서 USV를 향해 조준"),
    ("D. 통신", "RP-SMA 저손실 연장 케이블 3m", "LMR-195 급", 2, 12000, "손실 최소화를 위해 짧을수록 좋음"),
    ("D. 통신", "안테나 거치용 삼각대", "1.5m 이상", 1, 25000, "안테나 높이 = 통달거리"),
    ("D. 통신", "[선택] LoRa 모듈 E22-900T22S", "920MHz 대역 설정 필수", 2, 18000, "영상 제외 백업 텔레메트리"),
    ("D. 통신", "[선택] 920MHz 안테나 (SMA)", "3dBi", 2, 7000, "LoRa용"),

    ("E. RC 조종", "FlySky FS-i6 + FS-iA6B 수신기", "2.4GHz 6채널, iBUS 출력", 1, 85000,
     "★ 비상 정지·수동 탈출용 안전장치"),
    ("E. RC 조종", "AA 건전지 4개", "송신기용", 1, 5000, ""),

    ("F. 선체·기구", "PETG 필라멘트 1kg", "선체·데크 출력용 (흰색/주황 권장)", 3, 28000,
     "★ PLA 금지 - 햇빛·수온에 변형"),
    ("F. 선체·기구", "TPU 필라멘트 0.5kg", "쇼어 95A", 1, 25000, "방수 개스킷·범퍼"),
    ("F. 선체·기구", "방수 인클로저 IP65", "약 200×120×75mm ABS", 1, 18000, "전자부 수납"),
    ("F. 선체·기구", "방수 케이블 그랜드 PG7", "IP68", 6, 700, "배선 관통부"),
    ("F. 선체·기구", "알루미늄 각파이프 20×20×1.5t", "길이 1m", 2, 7000, "좌우 선체 연결 빔"),
    ("F. 선체·기구", "EPP/EPS 폼 블록", "예비 부력재", 2, 5000, "선체 내부 충전 - 침수해도 뜸"),
    ("F. 선체·기구", "M3 볼트·너트·와셔 세트", "10~30mm 혼합", 1, 12000, "스테인리스 권장"),
    ("F. 선체·기구", "M3 황동 인서트 너트 50개", "열삽입형", 1, 8000, "3D 출력물 체결"),
    ("F. 선체·기구", "중성 실리콘 실란트 (투명)", "습기 경화형", 1, 6000, "선체 이음부 방수"),
    ("F. 선체·기구", "부틸 방수 테이프", "폭 25mm", 1, 5000, "커넥터부 마감"),
    ("F. 선체·기구", "벨크로 스트랩 / 케이블 타이 세트", "", 1, 6000, "배터리 고정"),

    ("G. 공구·소모품", "온도조절형 납땜인두 세트", "60W 급", 1, 35000, "학교 보유 시 제외"),
    ("G. 공구·소모품", "열수축튜브 세트", "Ø2–10mm", 1, 5000, ""),
    ("G. 공구·소모품", "디지털 멀티미터", "전압/전류/도통", 1, 20000, "학교 보유 시 제외"),
    ("G. 공구·소모품", "육각·드라이버 공구세트", "", 1, 15000, ""),
    ("G. 공구·소모품", "디지털 주방저울 5kg", "추력 측정 시험용", 1, 12000, "지상 추력 시험에 사용"),
]

r = 5
for n, (grp, name, spec, qty, price, note) in enumerate(ITEMS, 1):
    ws.cell(row=r, column=1, value=n)
    ws.cell(row=r, column=2, value=grp)
    ws.cell(row=r, column=3, value=name)
    ws.cell(row=r, column=4, value=spec)
    ws.cell(row=r, column=5, value=qty)
    ws.cell(row=r, column=6, value=price)
    ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
    ws.cell(row=r, column=8, value=note)
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.font = BLACK
        cell.border = BOX
        cell.alignment = Alignment(vertical="center",
                                   wrap_text=(c in (4, 8)),
                                   horizontal="center" if c in (1, 5) else "left")
    ws.cell(row=r, column=6).number_format = "#,##0"
    ws.cell(row=r, column=7).number_format = "#,##0"
    if n % 2 == 0:
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = GFILL
    r += 1

last = r - 1
ws.cell(row=r, column=3, value="합계").font = BOLD
ws.cell(row=r, column=7, value=f"=SUM(G5:G{last})").font = BOLD
ws.cell(row=r, column=7).number_format = "#,##0"
ws.cell(row=r, column=7).fill = YFILL
for c in range(1, 9):
    ws.cell(row=r, column=c).border = BOX

r += 2
ws.cell(row=r, column=3, value="※ ★ 표시는 안전 또는 성능상 반드시 지켜야 하는 항목입니다.").font = SUB
ws.cell(row=r + 1, column=3,
        value="※ 3D 프린팅 출력물(선체·마운트·가드)은 필라멘트 비용에 포함되어 단가 0원으로 표기했습니다.").font = SUB

for col, w in zip("ABCDEFGH", [6, 14, 34, 42, 7, 11, 12, 40]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:H{last}"

# =====================================================================
# 시트 2 : 예산 요약
# =====================================================================
s2 = wb.create_sheet("예산요약")
s2["A1"] = "구분별 예산 요약"
s2["A1"].font = TITLE
for i, h in enumerate(["구분", "금액(원)", "비율"], 1):
    c = s2.cell(row=3, column=i, value=h)
    c.font = HEAD
    c.fill = HFILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center")

GROUPS = ["A. 추진계", "B. 제어·센서", "C. 전원", "D. 통신",
          "E. RC 조종", "F. 선체·기구", "G. 공구·소모품"]
rr = 4
for g in GROUPS:
    s2.cell(row=rr, column=1, value=g).font = BLACK
    s2.cell(row=rr, column=2,
            value=f'=SUMIF(BOM!$B$5:$B${last},A{rr},BOM!$G$5:$G${last})').font = BLACK
    s2.cell(row=rr, column=2).number_format = "#,##0"
    for c in range(1, 4):
        s2.cell(row=rr, column=c).border = BOX
    rr += 1
tot = rr
s2.cell(row=tot, column=1, value="총계").font = BOLD
s2.cell(row=tot, column=2, value=f"=SUM(B4:B{tot - 1})").font = BOLD
s2.cell(row=tot, column=2).number_format = "#,##0"
s2.cell(row=tot, column=2).fill = YFILL
for c in range(1, 4):
    s2.cell(row=tot, column=c).border = BOX
for i in range(4, tot):
    cell = s2.cell(row=i, column=3, value=f"=IFERROR(B{i}/$B${tot},0)")
    cell.number_format = "0.0%"
    cell.font = BLACK
    cell.border = BOX

s2.cell(row=tot + 2, column=1,
        value="※ 공구(G) 를 학교에서 이미 보유 중이라면 그만큼 차감됩니다.").font = SUB
s2.cell(row=tot + 3, column=1,
        value="※ 2호기부터는 공구·충전기·송신기·지상국 안테나가 재사용되어 비용이 크게 줄어듭니다.").font = SUB
for col, w in zip("ABC", [22, 16, 10]):
    s2.column_dimensions[col].width = w

# =====================================================================
# 시트 3 : 추진 성능 계산 (A2212 적합성 검증)
# =====================================================================
s3 = wb.create_sheet("추진성능계산")
s3["A1"] = "A2212 10T 1400KV 적합성 검증 · 추진 성능 계산기"
s3["A1"].font = TITLE
s3["A2"] = "파란 글씨 칸(B열)만 바꾸면 나머지가 자동으로 다시 계산됩니다."
s3["A2"].font = SUB


def block(row, title):
    c = s3.cell(row=row, column=1, value=title)
    c.font = HEAD
    c.fill = HFILL
    c.border = BOX
    for k in (2, 3, 4):
        s3.cell(row=row, column=k).fill = HFILL
        s3.cell(row=row, column=k).border = BOX


def line(row, label, value, unit, note="", inp=False, fmt="0.00"):
    s3.cell(row=row, column=1, value=label).font = BLACK
    c = s3.cell(row=row, column=2, value=value)
    c.font = BLUE if inp else BLACK
    c.number_format = fmt
    if inp:
        c.fill = YFILL
    s3.cell(row=row, column=3, value=unit).font = BLACK
    s3.cell(row=row, column=4, value=note).font = SUB
    for k in range(1, 5):
        s3.cell(row=row, column=k).border = BOX


block(4, "① 입력값 (노란 칸을 직접 바꿔 보세요)")
line(5,  "배터리 셀 수 (S)",            3,    "S",     "3S 고정 — 모터 사양 한계", True, "0")
line(6,  "셀 공칭 전압",                3.7,  "V",     "LiPo 표준", True)
line(7,  "모터 KV",                     1400, "rpm/V", "A2212 10T", True, "0")
line(8,  "모터 개수",                   2,    "개",    "좌현 + 우현", True, "0")
line(9,  "프로펠러 1개 최대 정지추력",   780,  "g",     '8×4.5" @3S 실측 참고값', True, "0")
line(10, "프로펠러 1개 최대 전류",       13,   "A",     '8×4.5" @3S 실측 참고값', True, "0")
line(11, "ESC 정격 전류",               30,   "A",     "BLHeli_S 30A", True, "0")
line(12, "순항 스로틀",                 0.55, "비율",  "실제 운용 시 55% 권장", True, "0%")
line(13, "선체 총중량(완성 상태)",       2.5,  "kg",    "선체+전자부+배터리", True, "0.0")
line(14, "배터리 용량",                 5000, "mAh",   "3S 5000mAh", True, "0")
line(15, "배터리 안전 사용률",           0.8,  "비율",  "완전 방전 금지", True, "0%")
line(16, "전자부 상시 소비전류",         0.6,  "A",     "ESP32+카메라+공유기 합", True, "0.00")
line(17, "실운용 보정계수",             0.7,  "비율",  "파도·바람·선회 손실 반영. 실측 후 조정", True, "0%")

block(19, "② 계산 결과")
s3["A20"] = "배터리 공칭 전압";      s3["B20"] = "=B5*B6";               s3["C20"] = "V"
s3["D20"] = "만충 시에는 4.2V/셀 = 12.6V"
s3["A21"] = "무부하 최대 회전수";    s3["B21"] = "=B7*B20";              s3["C21"] = "rpm"
s3["D21"] = "실부하에서는 이보다 낮음"
s3["A22"] = "최대 정지추력 (합계)";  s3["B22"] = "=B9*B8";               s3["C22"] = "g"
s3["D22"] = "두 모터 100% 출력 시"
s3["A23"] = "추력 대 중량비";        s3["B23"] = "=B22/(B13*1000)";      s3["C23"] = "배"
s3["D23"] = "0.3 이상이면 충분 (배는 항공기와 달리 1을 넘길 필요 없음)"
s3["A24"] = "순항 추력";             s3["B24"] = "=B22*B12^2";           s3["C24"] = "g"
s3["D24"] = "추력 ∝ 스로틀²"
s3["A25"] = "순항 시 모터 전류(합)"; s3["B25"] = "=B10*B8*B12^3";        s3["C25"] = "A"
s3["D25"] = "전류 ∝ 스로틀³ (근사)"
s3["A26"] = "순항 시 총 소비전류";   s3["B26"] = "=B25+B16";             s3["C26"] = "A"
s3["D26"] = "전자부 포함"
s3["A27"] = "이론 운용시간";         s3["B27"] = "=B14/1000/B26*60*B15"; s3["C27"] = "분"
s3["D27"] = "안전 사용률 반영, 이상적인 조건"
s3["A28"] = "예상 실운용 시간";      s3["B28"] = "=B27*B17";             s3["C28"] = "분"
s3["D28"] = "★ 실제로 기대할 수 있는 값"
s3["A29"] = "모터 1개당 최대 전력";  s3["B29"] = "=B20*B10";             s3["C29"] = "W"
s3["D29"] = "A2212 정격 180W"
s3["A30"] = "ESC 전류 여유율";       s3["B30"] = "=B11/B10";             s3["C30"] = "배"
s3["D30"] = "1.5배 이상 권장"

for row in range(20, 31):
    s3.cell(row=row, column=1).font = BLACK
    s3.cell(row=row, column=2).font = BLACK
    s3.cell(row=row, column=2).number_format = "#,##0.00"
    s3.cell(row=row, column=3).font = BLACK
    s3.cell(row=row, column=4).font = SUB
    for k in range(1, 5):
        s3.cell(row=row, column=k).border = BOX
s3["B21"].number_format = "#,##0"
s3["B22"].number_format = "#,##0"
s3["B27"].number_format = "#,##0.0"
s3["B28"].number_format = "#,##0.0"

block(32, "③ 판정")
JUDGE = [
    ("모터 정격 전력", '=IF(B29<=180,"적합 — 정격 180W 이내","부적합 — 프로펠러를 더 작게")'),
    ("ESC 여유율",     '=IF(B30>=1.5,"적합 — 30A ESC 여유 충분","주의 — 더 큰 ESC 필요")'),
    ("추력 대 중량비", '=IF(B23>=0.3,"적합 — 쌍동선 항주에 충분","부족 — 경량화 또는 프로펠러 확대")'),
    ("운용 시간",      '=IF(B28>=15,"양호 — 15분 이상 운용 가능","짧음 — 배터리 용량 증설 권장")'),
]
jr = 33
for label, f in JUDGE:
    s3.cell(row=jr, column=1, value=label).font = BLACK
    c = s3.cell(row=jr, column=2, value=f)
    c.font = BOLD
    s3.merge_cells(start_row=jr, start_column=2, end_row=jr, end_column=4)
    for k in range(1, 5):
        s3.cell(row=jr, column=k).border = BOX
    jr += 1

s3.cell(row=jr + 1, column=1,
        value='※ 정지추력·전류값은 A2212 1400KV + 8×4.5" 프로펠러 + 3S 조합의 공개 실측 참고치입니다.').font = SUB
s3.cell(row=jr + 2, column=1,
        value="※ 완성 후 주방저울로 실제 추력을 측정해 B9·B10 값을 우리 기체 실측치로 바꾸면 정확해집니다.").font = SUB

for col, w in zip("ABCD", [28, 14, 8, 52]):
    s3.column_dimensions[col].width = w

# =====================================================================
# 시트 4 : 통신 방식 비교
# =====================================================================
s4 = wb.create_sheet("통신방식비교")
s4["A1"] = "장거리 통신 확보 방안 비교"
s4["A1"].font = TITLE
s4["A2"] = ("실사용 거리는 수면 상태·주변 전파 혼잡도·안테나 높이에 따라 크게 달라집니다. "
            "아래는 개활 수면 기준 대략치입니다.")
s4["A2"].font = SUB

CH = ["방안", "구성", "예상 통달거리", "영상", "추가 비용", "난이도", "권장도"]
for i, h in enumerate(CH, 1):
    c = s4.cell(row=4, column=i, value=h)
    c.font = HEAD
    c.fill = HFILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")

ROWS = [
    ("0. 기본형 (비교 기준)", "미니 공유기 기본 안테나 + 노트북 내장 Wi-Fi",
     "80 ~ 150 m", "가능", "0원", "쉬움", "연습용"),
    ("1. 안테나 교체 + 높이 확보", "USV에 5dBi 안테나, 지상 안테나를 삼각대로 1.5m 이상 올림",
     "200 ~ 350 m", "가능", "약 3만원", "쉬움", "★★★ 1순위"),
    ("2. 지상국 지향성 안테나", "방안1 + 지상국에 14dBi 패널 안테나 + 고출력 USB 무선랜",
     "600 m ~ 1.2 km", "가능", "약 9만원", "보통", "★★★ 핵심"),
    ("3. 채널·대역폭 최적화", "2.4GHz 고정 · 20MHz 폭 고정 · 한산한 채널(1/6/11) 선택 · 5GHz 끄기",
     "방안2의 1.3배", "가능", "0원", "보통", "★★★ 반드시 적용"),
    ("4. LoRa 백업 텔레메트리", "920MHz E22 모듈 2개로 센서·명령 채널만 이중화 (영상 제외)",
     "1 ~ 3 km", "불가", "약 5만원", "어려움", "★★ 대회·장거리 시"),
    ("5. Wi-Fi 중계기(수상 부표)", "중간 지점에 부표형 중계 공유기 설치",
     "구간마다 +300 m", "가능", "약 5만원", "어려움", "★ 특수 상황"),
]
rr = 5
for row in ROWS:
    for i, v in enumerate(row, 1):
        c = s4.cell(row=rr, column=i, value=v)
        c.font = BLACK
        c.border = BOX
        c.alignment = Alignment(vertical="center", wrap_text=(i == 2),
                                horizontal="center" if i in (3, 4, 5, 6, 7) else "left")
    if rr % 2 == 1:
        for i in range(1, 8):
            s4.cell(row=rr, column=i).fill = GFILL
    rr += 1

s4.cell(row=rr + 1, column=1,
        value="권장 조합: 방안 1 + 2 + 3 을 함께 적용 (약 12만원 추가로 1km 급 확보)").font = BOLD
s4.cell(row=rr + 2, column=1,
        value="※ 국내에서 920MHz(917–923.5MHz) 대역은 비면허 소출력 무선기기용으로 개방되어 있습니다. "
              "LoRa 모듈은 반드시 이 대역 채널로 설정하고, 가능하면 국내 적합성평가(KC) 인증 제품을 사용하세요.").font = SUB
s4.cell(row=rr + 3, column=1,
        value="※ Wi-Fi 송신 출력을 규정 이상으로 높이는 개조(파워 부스터 등)는 전파법 위반입니다. "
              "출력이 아니라 '안테나 이득과 높이'로 거리를 버는 것이 합법적이고 효과도 큽니다.").font = SUB

for col, w in zip("ABCDEFG", [24, 46, 16, 8, 12, 9, 16]):
    s4.column_dimensions[col].width = w

# =====================================================================
import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "대전중앙중_USV_BOM.xlsx")
wb.save(out)
print("saved:", out)
